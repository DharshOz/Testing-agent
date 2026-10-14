import sys
from pathlib import Path

# Add project root to path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import csv
import hashlib
import json
import shutil
import pandas as pd
import yaml
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document
from snapshot import create_snapshot_for_case


def create_pdf(file_path: Path, title: str, lines: list[str]):
    file_path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(file_path), pagesize=letter)
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, 750, title)
    c.setFont("Helvetica", 10)
    y = 710
    for line in lines:
        c.drawString(50, y, line)
        y -= 18
        if y < 50:
            c.showPage()
            c.setFont("Helvetica", 10)
            y = 750
    c.save()


def create_docx(file_path: Path, title: str, paragraphs: list[str]):
    file_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(title, level=0)
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(str(file_path))


def create_txt(file_path: Path, text: str):
    file_path.parent.mkdir(parents=True, exist_ok=True)
    file_path.write_text(text, encoding="utf-8")


def generate_demo_issue():
    case_id = "case_demo_id_prefix"
    case_dir = Path("tests_agent/cases") / case_id
    files_dir = case_dir / "files"
    files_dir.mkdir(parents=True, exist_ok=True)

    print(f"Creating demonstration documents for {case_id}...")

    # 1. Application Form (has prefix: AADHAAR-8921-4321-9988)
    app_pdf = files_dir / "application_form.pdf"
    create_pdf(app_pdf, "CIVIC GRANT APPLICATION - DEMONSTRATION ISSUE", [
        "Applicant Name: Aarav Sharma",
        "Applicant ID: AADHAAR-8921-4321-9988",
        "Contact: +91-9876543210, aarav.sharma@example.gov.in",
        "Date of Submission: 2026-03-15",
        "Department: Rural Water Supply and Sanitation",
        "Issue Title: Solar Water Pump Infrastructure Project",
        "Amount Requested: INR 2,50,000",
        "Supporting References: SANCTION-ORD-2026-449",
        "Project Description: Solar water pump project demonstration ticket."
    ])

    # 2. Aadhaar Identity Proof (clean numeric only: 8921-4321-9988)
    id_txt = files_dir / "aadhaar_identity_proof.txt"
    create_txt(id_txt, """GOVERNMENT OF INDIA - UNIQUE IDENTIFICATION AUTHORITY
IDENTITY PROOF CARD (AADHAAR)
Name: Aarav Sharma
Father's Name: R. K. Sharma
DOB: 1985-06-12
Gender: Male
Aadhaar Number: 8921-4321-9988
Address: House No. 42, Ward 4, Greenfield Village, District North
Status: ACTIVE VERIFIED CITIZEN
""")

    # 3. Address Proof
    addr_docx = files_dir / "address_proof_utility_bill.docx"
    create_docx(addr_docx, "STATE ELECTRICITY DISTRIBUTION UTILITY", [
        "Consumer Name: Aarav Sharma",
        "Consumer ID: ELEC-99210-W4",
        "Service Address: House No. 42, Ward 4, Greenfield Village",
        "Verification: Verified residential utility bill."
    ])

    # 4. Sanction Order
    sanction_pdf = files_dir / "sanction_order_approval.pdf"
    create_pdf(sanction_pdf, "DEPARTMENT PRELIMINARY SANCTION ORDER", [
        "Sanction Order Reference: SANCTION-ORD-2026-449",
        "Date: 2026-03-01",
        "Grant Beneficiary: Aarav Sharma",
        "Approved Scheme: Solar Water Pump Infrastructure Project",
        "Sanctioned Ceiling Amount: INR 2,50,000"
    ])

    # 5. Community Consent Resolution
    consent_docx = files_dir / "community_consent_resolution.docx"
    create_docx(consent_docx, "GRAM SABHA COMMUNITY CONSENT RESOLUTION", [
        "Resolution No: GS-RES-2026-018",
        "Meeting Date: 2026-03-10",
        "Subject: Community Consent for Solar Water Pump",
        "Consent Declaration: We hereby confirm collective community consent.",
        "Signatures and Approvals:",
        "1. Sunita Devi - Signed",
        "2. Manoj Kumar Patel - Signed",
        "Attested by Panchayat Secretary on 2026-03-10."
    ])

    # Write input.json
    file_entries = []
    for f in [app_pdf, id_txt, addr_docx, sanction_pdf, consent_docx]:
        sha = hashlib.sha256(f.read_bytes()).hexdigest()
        file_entries.append({
            "filename": f.name,
            "path": f"files/{f.name}",
            "sha256": sha
        })

    input_data = {
        "ticket_id": f"TICK-{case_id.upper()}",
        "category": "grant_request",
        "files": file_entries
    }
    with open(case_dir / "input.json", "w", encoding="utf-8") as f:
        json.dump(input_data, f, indent=2)

    # Write meta.yaml
    meta_data = {
        "case_id": case_id,
        "data_type": "prefixed_identifier",
        "split": "tune",
        "must_pass": True
    }
    with open(case_dir / "meta.yaml", "w", encoding="utf-8") as f:
        yaml.safe_dump(meta_data, f)

    # Write custom expected.yaml based on the issue
    expected_data = {
        "extract_text": {
            "expected_files": [f["filename"] for f in file_entries]
        },
        "extract_fields": {
            "absent_fields": []
        },
        "check_documents": {
            "all_required_present": True,
            "missing_types": []
        },
        "validate_ticket": {
            "verdict": "valid",
            "judge_rubrics": [
                "Validator must recognize that AADHAAR-8921-4321-9988 and 8921-4321-9988 represent the exact same valid identity and not flag a false mismatch."
            ]
        },
        "generate_report": {
            "required_headings": ["Executive Summary", "Applicant & Submission Details", "Civic Request Details"],
            "applicant_name": "Aarav Sharma"
        }
    }
    with open(case_dir / "expected.yaml", "w", encoding="utf-8") as f:
        yaml.safe_dump(expected_data, f)

    # Run snapshot
    print(f"Creating golden snapshots for {case_id}...")
    create_snapshot_for_case(case_id)

    # Update issues/issues.csv and issues/issues.xlsx
    issues_csv = Path("issues/issues.csv")
    issues_xlsx = Path("issues/issues.xlsx")

    # Read existing
    rows = []
    if issues_csv.is_file():
        with open(issues_csv, "r", encoding="utf-8") as f:
            rows = list(csv.DictReader(f))

    # Check if demo row already exists
    demo_exists = any(r.get("test_case") == case_id for r in rows)
    if not demo_exists:
        new_id = len(rows) + 1
        rows.append({
            "id": str(new_id),
            "stage": "validate_ticket",
            "data_type": "prefixed_identifier",
            "issue": "Prefix-formatted Applicant IDs cause false cross-document validation failure",
            "reason": "Application form has 'AADHAAR-8921-4321-9988' while ID proof has '8921-4321-9988'. Validator prompt needs generic instruction to normalize common identity prefixes before matching.",
            "test_case": case_id,
            "expected_behavior": "Validator must recognize prefix equivalence and mark ticket valid with passing ID verification check.",
            "status": "open",
            "attempts": "0",
            "resolution_note": "",
            "linked_pr": ""
        })

        with open(issues_csv, "w", newline="", encoding="utf-8") as f:
            fieldnames = [
                "id", "stage", "data_type", "issue", "reason", "test_case",
                "expected_behavior", "status", "attempts", "resolution_note", "linked_pr"
            ]
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)

    # Save to Excel as well!
    df = pd.DataFrame(rows)
    df.to_excel(issues_xlsx, index=False)
    print(f"[OK] issues.csv and issues.xlsx synchronized with {len(rows)} issues.")
    print(f"[OK] Demonstration issue setup completed for {case_id}!")


if __name__ == "__main__":
    generate_demo_issue()
