import os
from pathlib import Path
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


def create_simple_pdf(file_path: Path, title: str, content_lines: list[str]):
    file_path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(file_path), pagesize=letter)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, 750, title)
    c.setFont("Helvetica", 11)
    y = 710
    for line in content_lines:
        c.drawString(50, y, line)
        y -= 20
        if y < 50:
            c.showPage()
            c.setFont("Helvetica", 11)
            y = 750
    c.save()


def create_simple_docx(file_path: Path, title: str, paragraphs: list[str]):
    file_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(title, level=0)
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(str(file_path))


def create_simple_txt(file_path: Path, content: str):
    file_path.parent.mkdir(parents=True, exist_ok=True)
    file_path.write_text(content, encoding="utf-8")


def generate_all_samples():
    base_dir = Path("sample_data")
    base_dir.mkdir(parents=True, exist_ok=True)

    # =========================================================================
    # TICKET 1: Fully Valid (grant_request)
    # =========================================================================
    t1_dir = base_dir / "ticket_1_valid"
    t1_dir.mkdir(parents=True, exist_ok=True)

    # 1. Application form (PDF)
    create_simple_pdf(
        t1_dir / "application_form.pdf",
        "CIVIC GRANT APPLICATION FORM",
        [
            "Applicant Name: Aarav Sharma",
            "Applicant ID: AADHAAR-8921-4321-9988",
            "Contact: +91-9876543210, aarav.sharma@example.gov.in",
            "Date of Submission: 2026-03-15",
            "Department: Rural Water Supply and Sanitation",
            "Issue Title: Installation of Community Solar Water Pump in Ward 4",
            "Amount Requested: INR 2,50,000",
            "Supporting References: SANCTION-ORD-2026-449",
            "",
            "Project Description:",
            "The community of Ward 4 has been experiencing recurrent drinking water shortages during peak summer.",
            "This application requests an infrastructure grant of INR 2,50,000 for drilling and installing a high-efficiency",
            "community solar-powered submersible water pump at the Central Civic Ground.",
            "All preliminary engineering clearances and local Gram Sabha approvals have been completed."
        ]
    )

    # 2. Identity Proof (TXT)
    create_simple_txt(
        t1_dir / "aadhaar_identity_proof.txt",
        """GOVERNMENT OF INDIA - UNIQUE IDENTIFICATION AUTHORITY
IDENTITY PROOF CARD (AADHAAR)
Name: Aarav Sharma
Father's Name: R. K. Sharma
DOB: 1985-06-12
Gender: Male
Aadhaar Number: AADHAAR-8921-4321-9988
Address: House No. 42, Ward 4, Greenfield Village, District North
Status: ACTIVE VERIFIED CITIZEN
"""
    )

    # 3. Address Proof (DOCX)
    create_simple_docx(
        t1_dir / "address_proof_utility_bill.docx",
        "STATE ELECTRICITY DISTRIBUTION UTILITY - RESIDENCE PROOF",
        [
            "Consumer Name: Aarav Sharma",
            "Consumer ID: ELEC-99210-W4",
            "Service Address: House No. 42, Ward 4, Greenfield Village",
            "Billing Cycle: February - March 2026",
            "Bill Status: Paid in Full",
            "Verification: This document certifies valid domestic residence proof for Aarav Sharma at Ward 4."
        ]
    )

    # 4. Sanction Document (PDF)
    create_simple_pdf(
        t1_dir / "sanction_order_approval.pdf",
        "DEPARTMENT OF RURAL WATER SUPPLY - PRELIMINARY SANCTION ORDER",
        [
            "Sanction Order Reference: SANCTION-ORD-2026-449",
            "Date: 2026-03-01",
            "Grant Beneficiary: Aarav Sharma (Lead Community Representative, Ward 4)",
            "Approved Scheme: Installation of Community Solar Water Pump in Ward 4",
            "Sanctioned Ceiling Amount: INR 2,50,000",
            "Issuing Officer: Executive Engineer, Department of Rural Water Supply",
            "Terms: The sanction is subject to formal community consent resolution and identity verification."
        ]
    )

    # 5. Community Consent Document (DOCX)
    create_simple_docx(
        t1_dir / "community_consent_resolution.docx",
        "GRAM SABHA FORMAL COMMUNITY CONSENT RESOLUTION",
        [
            "Resolution No: GS-RES-2026-018",
            "Meeting Date: 2026-03-10",
            "Location: Ward 4 Community Center Hall",
            "Subject: Formal Public Consent for Installation of Community Solar Water Pump",
            "",
            "Consent Declaration:",
            "We, the undersigned permanent residents of Ward 4, hereby solemnly confirm our collective consent",
            "for the proposed solar water pump installation submitted by Shri Aarav Sharma (Lead Representative).",
            "We agree to the designated site at the Central Civic Ground and pledge community maintenance support.",
            "",
            "Resident Signatures and Approvals:",
            "1. Sunita Devi (Resident #12) - Signed",
            "2. Manoj Kumar Patel (Resident #15) - Signed",
            "3. Lakshmi Narayanan (Resident #22) - Signed",
            "4. Priya Sen (Resident #07) - Signed",
            "5. Gurpreet Singh (Resident #31) - Signed",
            "6. Mohammed Rizwan (Resident #19) - Signed",
            "7. Anita Deshmukh (Resident #44) - Signed",
            "8. Ramesh Chand (Resident #05) - Signed",
            "",
            "Attested by: Panchayat Secretary & Ward 4 Civic Officer on 2026-03-10."
        ]
    )

    # =========================================================================
    # TICKET 2: Missing Consent Document (grant_request)
    # =========================================================================
    t2_dir = base_dir / "ticket_2_missing_consent"
    t2_dir.mkdir(parents=True, exist_ok=True)

    create_simple_pdf(
        t2_dir / "application_form.pdf",
        "GRANT APPLICATION - STREET PAVING REPAIR",
        [
            "Applicant Name: Meera Joshi",
            "Applicant ID: AADHAAR-3344-5566-7788",
            "Contact: +91-9123456780, meera.j@sample.org",
            "Date: 2026-04-02",
            "Department: Municipal Infrastructure Board",
            "Issue Title: Street Paving and Drainage Repair in South Block",
            "Amount Requested: INR 1,80,000",
            "Description: Application for repairing damaged roadway and laying concrete tiles."
        ]
    )

    create_simple_txt(
        t2_dir / "aadhaar_identity_proof.txt",
        """GOVERNMENT OF INDIA - AADHAAR IDENTITY CARD
Name: Meera Joshi
DOB: 1991-11-20
Aadhaar Number: AADHAAR-3344-5566-7788
Address: Flat 102, South Block Apartments, City Zone 3
Status: ACTIVE
"""
    )

    create_simple_docx(
        t2_dir / "address_proof_utility_bill.docx",
        "CITY WATER BOARD RESIDENTIAL PROOF",
        [
            "Consumer Name: Meera Joshi",
            "Address: Flat 102, South Block Apartments",
            "Statement Date: March 2026",
            "Document: Valid Residential Address Proof"
        ]
    )

    create_simple_pdf(
        t2_dir / "sanction_approval_clearance.pdf",
        "MUNICIPAL INFRASTRUCTURE BOARD - CONDITIONAL SANCTION",
        [
            "Sanction Ref: SANCTION-MIB-2026-901",
            "Applicant: Meera Joshi",
            "Scope: Street Paving South Block",
            "Amount: INR 1,80,000",
            "Condition: Mandatory submission of community consent signatures before work order."
        ]
    )
    # NOTE: community_consent_document is intentionally omitted!

    # =========================================================================
    # TICKET 3: Inconsistent Names Across Documents (grant_request)
    # =========================================================================
    t3_dir = base_dir / "ticket_3_inconsistent_names"
    t3_dir.mkdir(parents=True, exist_ok=True)

    create_simple_pdf(
        t3_dir / "application_form.pdf",
        "COMMUNITY GRANT REQUEST FORM",
        [
            "Applicant Name: Vikram Malhotra",
            "Applicant ID: AADHAAR-5544-3322-1100",
            "Contact: +91-9988776655, vikram.m@domain.com",
            "Date: 2026-05-10",
            "Department: Community Welfare Department",
            "Issue Title: Library Reading Room Expansion",
            "Amount Requested: INR 3,00,000",
            "Description: Grant for expanding the public library reading room with 20 new desks."
        ]
    )

    # Inconsistent ID proof (name says "Rajesh Kumar Verma", ID is completely different)
    create_simple_txt(
        t3_dir / "aadhaar_identity_proof.txt",
        """GOVERNMENT OF INDIA - AADHAAR IDENTITY CARD
Name: Rajesh Kumar Verma
DOB: 1978-02-14
Aadhaar Number: AADHAAR-9988-7766-5544
Address: Sector 18, Block B, Industrial Area
Status: ACTIVE
"""
    )

    create_simple_docx(
        t3_dir / "address_proof_utility_bill.docx",
        "MUNICIPAL UTILITY RESIDENCE PROOF",
        [
            "Consumer Name: Dinesh Chawla",
            "Address: Plot 55, Sector 9",
            "Description: Utility verification document."
        ]
    )

    create_simple_pdf(
        t3_dir / "sanction_document_letter.pdf",
        "ADMINISTRATIVE SANCTION CLEARANCE",
        [
            "Sanction Order: SANCTION-CWD-2026-312",
            "Recipient Name: Suresh Oberoi",
            "Sanctioned Project: Solar Streetlights (Does not match Library Expansion)",
            "Amount: INR 1,20,000"
        ]
    )

    create_simple_docx(
        t3_dir / "community_consent_document.docx",
        "COMMUNITY CONSENT RECORD",
        [
            "Resolution Date: 2026-01-05",
            "Project: Installation of Park Benches for Senior Citizens",
            "Lead Signee: Anil Kapoor",
            "Signatures: 4 residents signed"
        ]
    )

    print("Synthetic sample data generated successfully in sample_data/ directory.")


if __name__ == "__main__":
    generate_all_samples()
