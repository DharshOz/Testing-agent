import sys
from pathlib import Path

# Add project root to path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import hashlib
import json
import os
import shutil

from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document
import yaml

from snapshot import create_snapshot_for_case


def create_pdf(file_path: Path, title: str, lines: list[str]):
    file_path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(file_path), pagesize=letter)
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, 750, title)
    c.setFont("Helvetica", 10)
    y = 710
    for line in lines:
        c.drawString(50, y, line)
        y -= 18
        if y < 50:
            c.showPage()
            c.setFont("Helvetica", 10)
            y = 750
    c.save()


def create_docx(file_path: Path, title: str, paragraphs: list[str]):
    file_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(title, level=0)
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(str(file_path))


def create_txt(file_path: Path, text: str):
    file_path.parent.mkdir(parents=True, exist_ok=True)
    file_path.write_text(text, encoding="utf-8")


def write_case_bundle(case_dir: Path, case_id: str, files_map: dict[str, Path], meta: dict, expected: dict, category: str = "grant_request"):
    case_dir.mkdir(parents=True, exist_ok=True)
    files_target_dir = case_dir / "files"
    files_target_dir.mkdir(parents=True, exist_ok=True)

    file_entries = []
    for rel_name, src_path in files_map.items():
        dst_path = files_target_dir / rel_name
        shutil.copy(src_path, dst_path)
        sha = hashlib.sha256(dst_path.read_bytes()).hexdigest()
        file_entries.append({
            "filename": rel_name,
            "path": f"files/{rel_name}",
            "sha256": sha
        })

    # input.json
    input_json = {
        "ticket_id": f"TICK-{case_id.upper()}",
        "category": category,
        "files": file_entries
    }
    with open(case_dir / "input.json", "w", encoding="utf-8") as f:
        json.dump(input_json, f, indent=2)

    # meta.yaml
    with open(case_dir / "meta.yaml", "w", encoding="utf-8") as f:
        yaml.safe_dump(meta, f)

    # expected.yaml
    with open(case_dir / "expected.yaml", "w", encoding="utf-8") as f:
        yaml.safe_dump(expected, f)


def build_all_seed_cases():
    scratch_dir = Path("scratch_cases_tmp")
    scratch_dir.mkdir(parents=True, exist_ok=True)
    cases_base = Path("tests_agent/cases")
    cases_base.mkdir(parents=True, exist_ok=True)

    print("Building reusable documents...")
    # Base valid docs
    app_form_pdf = scratch_dir / "application_form.pdf"
    create_pdf(app_form_pdf, "CIVIC GRANT APPLICATION FORM", [
        "Applicant Name: Aarav Sharma",
        "Applicant ID: AADHAAR-8921-4321-9988",
        "Contact: +91-9876543210, aarav.sharma@example.gov.in",
        "Date of Submission: 2026-03-15",
        "Department: Rural Water Supply and Sanitation",
        "Issue Title: Installation of Community Solar Water Pump in Ward 4",
        "Amount Requested: INR 2,50,000",
        "Supporting References: SANCTION-ORD-2026-449",
        "Project Description: Infrastructure grant for drilling and installing community solar water pump."
    ])

    id_proof_txt = scratch_dir / "aadhaar_identity_proof.txt"
    create_txt(id_proof_txt, """GOVERNMENT OF INDIA - UNIQUE IDENTIFICATION AUTHORITY
IDENTITY PROOF CARD (AADHAAR)
Name: Aarav Sharma
Father's Name: R. K. Sharma
DOB: 1985-06-12
Gender: Male
Aadhaar Number: AADHAAR-8921-4321-9988
Address: House No. 42, Ward 4, Greenfield Village, District North
Status: ACTIVE VERIFIED CITIZEN
""")

    addr_proof_docx = scratch_dir / "address_proof_utility_bill.docx"
    create_docx(addr_proof_docx, "STATE ELECTRICITY DISTRIBUTION UTILITY - RESIDENCE PROOF", [
        "Consumer Name: Aarav Sharma",
        "Consumer ID: ELEC-99210-W4",
        "Service Address: House No. 42, Ward 4, Greenfield Village",
        "Verification: Domestic residence address proof for Aarav Sharma."
    ])

    sanction_pdf = scratch_dir / "sanction_order_approval.pdf"
    create_pdf(sanction_pdf, "DEPARTMENT OF RURAL WATER SUPPLY - PRELIMINARY SANCTION ORDER", [
        "Sanction Order Reference: SANCTION-ORD-2026-449",
        "Date: 2026-03-01",
        "Grant Beneficiary: Aarav Sharma",
        "Approved Scheme: Installation of Community Solar Water Pump in Ward 4",
        "Sanctioned Ceiling Amount: INR 2,50,000"
    ])

    consent_docx = scratch_dir / "community_consent_resolution.docx"
    create_docx(consent_docx, "GRAM SABHA FORMAL COMMUNITY CONSENT RESOLUTION", [
        "Resolution No: GS-RES-2026-018",
        "Meeting Date: 2026-03-10",
        "Subject: Formal Public Consent for Installation of Community Solar Water Pump",
        "Consent Declaration: We, the undersigned permanent residents of Ward 4, hereby confirm collective consent.",
        "Signatures and Approvals:",
        "1. Sunita Devi (Resident #12) - Signed",
        "2. Manoj Kumar Patel (Resident #15) - Signed",
        "3. Lakshmi Narayanan (Resident #22) - Signed",
        "Attested by Panchayat Secretary on 2026-03-10."
    ])

    # 1. case_01_valid_grant (tune, must_pass: true)
    print("Generating case_01_valid_grant...")
    write_case_bundle(
        cases_base / "case_01_valid_grant",
        "case_01_valid_grant",
        {
            "application_form.pdf": app_form_pdf,
            "aadhaar_identity_proof.txt": id_proof_txt,
            "address_proof_utility_bill.docx": addr_proof_docx,
            "sanction_order_approval.pdf": sanction_pdf,
            "community_consent_resolution.docx": consent_docx
        },
        meta={"case_id": "case_01_valid_grant", "data_type": "standard_grant", "split": "tune", "must_pass": True},
        expected={
            "extract_text": {"expected_files": ["application_form.pdf", "aadhaar_identity_proof.txt", "address_proof_utility_bill.docx", "sanction_order_approval.pdf", "community_consent_resolution.docx"]},
            "extract_fields": {"absent_fields": []},
            "check_documents": {"all_required_present": True, "missing_types": []},
            "validate_ticket": {"verdict": "valid"},
            "generate_report": {"required_headings": ["Executive Summary", "Applicant & Submission Details", "Civic Request Details"], "applicant_name": "Aarav Sharma"}
        }
    )

    # 2. case_02_missing_consent (tune, must_pass: true)
    print("Generating case_02_missing_consent...")
    write_case_bundle(
        cases_base / "case_02_missing_consent",
        "case_02_missing_consent",
        {
            "application_form.pdf": app_form_pdf,
            "aadhaar_identity_proof.txt": id_proof_txt,
            "address_proof_utility_bill.docx": addr_proof_docx,
            "sanction_order_approval.pdf": sanction_pdf
        },
        meta={"case_id": "case_02_missing_consent", "data_type": "missing_doc", "split": "tune", "must_pass": True},
        expected={
            "check_documents": {"all_required_present": False, "missing_types": ["community_consent_document"]},
            "validate_ticket": {"verdict": "invalid"}
        }
    )

    # 3. case_03_missing_sanction (tune, must_pass: true)
    print("Generating case_03_missing_sanction...")
    write_case_bundle(
        cases_base / "case_03_missing_sanction",
        "case_03_missing_sanction",
        {
            "application_form.pdf": app_form_pdf,
            "aadhaar_identity_proof.txt": id_proof_txt,
            "address_proof_utility_bill.docx": addr_proof_docx,
            "community_consent_resolution.docx": consent_docx
        },
        meta={"case_id": "case_03_missing_sanction", "data_type": "missing_doc", "split": "tune", "must_pass": True},
        expected={
            "check_documents": {"all_required_present": False, "missing_types": ["sanction_document"]},
            "validate_ticket": {"verdict": "invalid"}
        }
    )

    # 4. case_04_inconsistent_names (tune, must_pass: true)
    print("Generating case_04_inconsistent_names...")
    mismatch_id_txt = scratch_dir / "mismatch_aadhaar.txt"
    create_txt(mismatch_id_txt, """GOVERNMENT OF INDIA - AADHAAR CARD
Name: Rajesh Kumar Verma
Aadhaar Number: AADHAAR-9988-7766-5544
Address: Sector 18, Block B, Industrial Area
Status: ACTIVE
""")
    mismatch_sanction_pdf = scratch_dir / "mismatch_sanction.pdf"
    create_pdf(mismatch_sanction_pdf, "DEPARTMENT SANCTION ORDER", [
        "Sanction Reference: SANCTION-CWD-2026-312",
        "Recipient Name: Suresh Oberoi",
        "Amount: INR 1,20,000"
    ])
    write_case_bundle(
        cases_base / "case_04_inconsistent_names",
        "case_04_inconsistent_names",
        {
            "application_form.pdf": app_form_pdf,
            "aadhaar_identity_proof.txt": mismatch_id_txt,
            "address_proof_utility_bill.docx": addr_proof_docx,
            "sanction_order_approval.pdf": mismatch_sanction_pdf,
            "community_consent_resolution.docx": consent_docx
        },
        meta={"case_id": "case_04_inconsistent_names", "data_type": "cross_doc_inconsistency", "split": "tune", "must_pass": True},
        expected={
            "check_documents": {"all_required_present": True, "missing_types": []},
            "validate_ticket": {"verdict": "invalid"}
        }
    )

    # 5. case_05_scanned_empty_pdf (tune, must_pass: true)
    print("Generating case_05_scanned_empty_pdf...")
    empty_pdf = scratch_dir / "empty_scanned_id.pdf"
    create_pdf(empty_pdf, "", [])  # totally empty canvas
    write_case_bundle(
        cases_base / "case_05_scanned_empty_pdf",
        "case_05_scanned_empty_pdf",
        {
            "application_form.pdf": app_form_pdf,
            "empty_scanned_id.pdf": empty_pdf,
            "address_proof_utility_bill.docx": addr_proof_docx,
            "sanction_order_approval.pdf": sanction_pdf,
            "community_consent_resolution.docx": consent_docx
        },
        meta={"case_id": "case_05_scanned_empty_pdf", "data_type": "scanned_empty", "split": "tune", "must_pass": True},
        expected={
            "extract_text": {"allow_empty_files": ["empty_scanned_id.pdf"]},
            "check_documents": {"empty_files": ["empty_scanned_id.pdf"]}
        }
    )

    # 6. case_06_duplicate_file (tune, must_pass: true)
    print("Generating case_06_duplicate_file...")
    write_case_bundle(
        cases_base / "case_06_duplicate_file",
        "case_06_duplicate_file",
        {
            "application_form.pdf": app_form_pdf,
            "application_form_copy.pdf": app_form_pdf,
            "aadhaar_identity_proof.txt": id_proof_txt,
            "address_proof_utility_bill.docx": addr_proof_docx,
            "sanction_order_approval.pdf": sanction_pdf,
            "community_consent_resolution.docx": consent_docx
        },
        meta={"case_id": "case_06_duplicate_file", "data_type": "duplicate_upload", "split": "tune", "must_pass": True},
        expected={
            "check_documents": {"duplicate_files": ["application_form_copy.pdf"]}
        }
    )

    # 7. case_07_informal_consent (heldout, must_pass: false)
    print("Generating case_07_informal_consent (heldout)...")
    informal_consent_docx = scratch_dir / "informal_note.docx"
    create_docx(informal_consent_docx, "COMMUNITY NOTE", [
        "To Whom It May Concern: Some people around here think the water pump is okay.",
        "That's all."
    ])
    write_case_bundle(
        cases_base / "case_07_informal_consent",
        "case_07_informal_consent",
        {
            "application_form.pdf": app_form_pdf,
            "aadhaar_identity_proof.txt": id_proof_txt,
            "address_proof_utility_bill.docx": addr_proof_docx,
            "sanction_order_approval.pdf": sanction_pdf,
            "community_consent_note.docx": informal_consent_docx
        },
        meta={"case_id": "case_07_informal_consent", "data_type": "informal_consent", "split": "heldout", "must_pass": False},
        expected={
            "validate_ticket": {"verdict": "invalid"}
        }
    )

    # 8. case_08_hallucination_trap (heldout, must_pass: false)
    print("Generating case_08_hallucination_trap (heldout)...")
    trap_app_pdf = scratch_dir / "trap_application.pdf"
    create_pdf(trap_app_pdf, "GRANT REQUEST SUMMARY", [
        "Applicant Name: Aarav Sharma",
        "Applicant ID: AADHAAR-8921-4321-9988",
        "Contact: +91-9876543210",
        "Date of Submission: 2026-03-15",
        "Issue Title: Public Ground Beautification",
        "Issue Description: General landscaping inquiry."
    ])
    write_case_bundle(
        cases_base / "case_08_hallucination_trap",
        "case_08_hallucination_trap",
        {
            "application_form.pdf": trap_app_pdf,
            "aadhaar_identity_proof.txt": id_proof_txt,
            "address_proof_utility_bill.docx": addr_proof_docx,
            "sanction_order_approval.pdf": sanction_pdf,
            "community_consent_resolution.docx": consent_docx
        },
        meta={"case_id": "case_08_hallucination_trap", "data_type": "hallucination_trap", "split": "heldout", "must_pass": False},
        expected={
            "extract_fields": {"absent_fields": ["amount_requested", "department"]}
        }
    )

    # Clean temporary scratch files
    shutil.rmtree(scratch_dir, ignore_errors=True)

    # Snapshot all 8 cases
    print("\nRunning initial snapshots for all 8 cases...")
    for i in range(1, 9):
        case_name = f"case_{i:02d}"
        matched = [p.name for p in cases_base.iterdir() if p.name.startswith(case_name)]
        if matched:
            print(f"Snapshotting {matched[0]}...")
            create_snapshot_for_case(matched[0])

    print("[OK] All 8 seed cases generated and snapshotted successfully.")


if __name__ == "__main__":
    build_all_seed_cases()
