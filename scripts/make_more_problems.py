import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import csv
import hashlib
import json
import pandas as pd
import yaml
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document
from snapshot import create_snapshot_for_case


def create_pdf(file_path: Path, title: str, lines: list[str]):
    file_path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(file_path), pagesize=letter)
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, 750, title)
    c.setFont("Helvetica", 10)
    y = 710
    for line in lines:
        c.drawString(50, y, line)
        y -= 18
        if y < 50:
            c.showPage()
            c.setFont("Helvetica", 10)
            y = 750
    c.save()


def create_docx(file_path: Path, title: str, paragraphs: list[str]):
    file_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(title, level=0)
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(str(file_path))


def create_txt(file_path: Path, text: str):
    file_path.parent.mkdir(parents=True, exist_ok=True)
    file_path.write_text(text, encoding="utf-8")


def generate_currency_problem():
    case_id = "case_demo_currency_format"
    case_dir = Path("tests_agent/cases") / case_id
    files_dir = case_dir / "files"
    files_dir.mkdir(parents=True, exist_ok=True)

    app_pdf = files_dir / "application_form.pdf"
    create_pdf(app_pdf, "CIVIC INFRASTRUCTURE GRANT REQUEST", [
        "Applicant Name: Sunita Rao",
        "Applicant ID: AADHAAR-1122-3344-5566",
        "Contact: +91-9811223344",
        "Date of Submission: 2026-03-20",
        "Department: Municipal Works Board",
        "Issue Title: Rainwater Harvesting Facility Installation",
        "Amount Requested: Rs. 3,75,000/- only",
        "Supporting References: SANCTION-MWB-2026-881",
        "Description: Financial assistance for community rainwater recharge tank."
    ])

    id_txt = files_dir / "aadhaar_identity_proof.txt"
    create_txt(id_txt, """GOVERNMENT OF INDIA - AADHAAR CARD
Name: Sunita Rao
Aadhaar Number: AADHAAR-1122-3344-5566
Address: Plot 9, Lake View Enclave
Status: ACTIVE
""")

    addr_docx = files_dir / "address_proof_utility_bill.docx"
    create_docx(addr_docx, "ELECTRICITY RESIDENCE PROOF", [
        "Consumer Name: Sunita Rao",
        "Address: Plot 9, Lake View Enclave",
        "Status: Verified Domestic Consumer"
    ])

    sanction_pdf = files_dir / "sanction_order_approval.pdf"
    create_pdf(sanction_pdf, "MUNICIPAL WORKS BOARD - SANCTION CLEARANCE", [
        "Sanction Order: SANCTION-MWB-2026-881",
        "Beneficiary: Sunita Rao",
        "Approved Amount: INR 375000",
        "Scheme: Rainwater Harvesting Facility Installation"
    ])

    consent_docx = files_dir / "community_consent_resolution.docx"
    create_docx(consent_docx, "RESIDENTS WELFARE ASSOCIATION CONSENT", [
        "Resolution Date: 2026-03-15",
        "Subject: Public Consent for Rainwater Recharge Tank",
        "Signatures: 10 Resident Signatories confirmed consent."
    ])

    file_entries = []
    for f in [app_pdf, id_txt, addr_docx, sanction_pdf, consent_docx]:
        sha = hashlib.sha256(f.read_bytes()).hexdigest()
        file_entries.append({"filename": f.name, "path": f"files/{f.name}", "sha256": sha})

    with open(case_dir / "input.json", "w", encoding="utf-8") as fj:
        json.dump({"ticket_id": f"TICK-{case_id.upper()}", "category": "grant_request", "files": file_entries}, fj, indent=2)

    with open(case_dir / "meta.yaml", "w", encoding="utf-8") as fm:
        yaml.safe_dump({"case_id": case_id, "data_type": "currency_formatting", "split": "tune", "must_pass": True}, fm)

    with open(case_dir / "expected.yaml", "w", encoding="utf-8") as fe:
        yaml.safe_dump({
            "extract_text": {"expected_files": [f["filename"] for f in file_entries]},
            "extract_fields": {"absent_fields": []},
            "check_documents": {"all_required_present": True, "missing_types": []},
            "validate_ticket": {
                "verdict": "valid",
                "judge_rubrics": ["Validator must recognize that 'Rs. 3,75,000/- only' and 'INR 375000' represent the same financial amount and not flag an amount discrepancy."]
            },
            "generate_report": {"required_headings": ["Executive Summary", "Applicant & Submission Details"]}
        }, fe)

    create_snapshot_for_case(case_id)
    print(f"[OK] Generated {case_id}")


def generate_date_order_problem():
    case_id = "case_demo_date_order"
    case_dir = Path("tests_agent/cases") / case_id
    files_dir = case_dir / "files"
    files_dir.mkdir(parents=True, exist_ok=True)

    app_pdf = files_dir / "application_form.pdf"
    create_pdf(app_pdf, "SAME-DAY GRANT APPLICATION", [
        "Applicant Name: Devendra Malik",
        "Applicant ID: AADHAAR-7788-9900-1122",
        "Contact: +91-9765432100",
        "Date of Submission: 2026-04-10",
        "Department: Urban Greenery Development",
        "Issue Title: Miyawaki Mini-Forest Plantation in Sector 12",
        "Amount Requested: INR 1,50,000",
        "Supporting References: SANCTION-UGD-2026-104",
        "Description: Urban afforestation proposal."
    ])

    id_txt = files_dir / "aadhaar_identity_proof.txt"
    create_txt(id_txt, """GOVERNMENT OF INDIA - AADHAAR CARD
Name: Devendra Malik
Aadhaar Number: AADHAAR-7788-9900-1122
Address: Block C-4, Sector 12
Status: ACTIVE
""")

    addr_docx = files_dir / "address_proof_utility_bill.docx"
    create_docx(addr_docx, "DOMESTIC ELECTRICITY BILL PROOF", [
        "Consumer: Devendra Malik",
        "Address: Block C-4, Sector 12",
        "Verification: Verified Domestic Residence"
    ])

    sanction_pdf = files_dir / "sanction_order_approval.pdf"
    create_pdf(sanction_pdf, "URBAN GREENERY DEVELOPMENT - ADMINISTRATIVE SANCTION", [
        "Sanction Reference: SANCTION-UGD-2026-104",
        "Date of Order: 2026-04-10",
        "Beneficiary: Devendra Malik",
        "Sanctioned Amount: INR 1,50,000"
    ])

    consent_docx = files_dir / "community_consent_resolution.docx"
    create_docx(consent_docx, "COMMUNITY CONSENT RESOLUTION", [
        "Meeting Date: 2026-04-08",
        "Subject: Public Consent for Mini-Forest Plantation",
        "Signatures: 12 neighborhood residents signed."
    ])

    file_entries = []
    for f in [app_pdf, id_txt, addr_docx, sanction_pdf, consent_docx]:
        sha = hashlib.sha256(f.read_bytes()).hexdigest()
        file_entries.append({"filename": f.name, "path": f"files/{f.name}", "sha256": sha})

    with open(case_dir / "input.json", "w", encoding="utf-8") as fj:
        json.dump({"ticket_id": f"TICK-{case_id.upper()}", "category": "grant_request", "files": file_entries}, fj, indent=2)

    with open(case_dir / "meta.yaml", "w", encoding="utf-8") as fm:
        yaml.safe_dump({"case_id": case_id, "data_type": "same_day_dates", "split": "tune", "must_pass": True}, fm)

    with open(case_dir / "expected.yaml", "w", encoding="utf-8") as fe:
        yaml.safe_dump({
            "extract_text": {"expected_files": [f["filename"] for f in file_entries]},
            "extract_fields": {"absent_fields": []},
            "check_documents": {"all_required_present": True, "missing_types": []},
            "validate_ticket": {
                "verdict": "valid",
                "judge_rubrics": ["Validator must accept that administrative sanction issued on the same date as application submission is valid and does not violate timeline integrity."]
            },
            "generate_report": {"required_headings": ["Executive Summary", "Applicant & Submission Details"]}
        }, fe)

    create_snapshot_for_case(case_id)
    print(f"[OK] Generated {case_id}")


def update_excel_and_csv():
    issues_csv = Path("issues/issues.csv")
    issues_xlsx = Path("issues/issues.xlsx")

    rows = [
        {
            "id": "1",
            "stage": "check_documents",
            "data_type": "missing_doc",
            "issue": "Missing community consent document not flagged early",
            "reason": "Keyword match erroneously identified generic resolution text as consent",
            "test_case": "case_02_missing_consent",
            "expected_behavior": "Check documents stage must flag community consent as missing when no valid consent document is present",
            "status": "fixed",
            "attempts": "1",
            "resolution_note": "Prompt and keyword rules updated to ensure formal consent resolution requirement is strictly matched",
            "linked_pr": "PR-101"
        },
        {
            "id": "2",
            "stage": "extract_fields",
            "data_type": "hallucination_trap",
            "issue": "Optional fields returned non-null when unstated in document",
            "reason": "Model guessed department from context instead of returning null when absent",
            "test_case": "case_08_hallucination_trap",
            "expected_behavior": "Model must return null for unstated fields without hallucinating",
            "status": "open",
            "attempts": "0",
            "resolution_note": "",
            "linked_pr": ""
        },
        {
            "id": "3",
            "stage": "validate_ticket",
            "data_type": "cross_doc_inconsistency",
            "issue": "Cross-document name mismatch not highlighted in checks list",
            "reason": "Subtle name variation was overlooked during validation evaluation",
            "test_case": "case_04_inconsistent_names",
            "expected_behavior": "Validator must check applicant name across all documents and report contradiction if mismatch exists",
            "status": "open",
            "attempts": "0",
            "resolution_note": "",
            "linked_pr": ""
        },
        {
            "id": "4",
            "stage": "validate_ticket",
            "data_type": "prefixed_identifier",
            "issue": "Prefix-formatted Applicant IDs cause false cross-document validation failure",
            "reason": "Application form has 'AADHAAR-8921-4321-9988' while ID proof has '8921-4321-9988'. Validator prompt needs instruction to normalize common identity prefixes before matching.",
            "test_case": "case_demo_id_prefix",
            "expected_behavior": "Validator must recognize prefix equivalence and mark ticket valid with passing ID verification check.",
            "status": "open",
            "attempts": "0",
            "resolution_note": "",
            "linked_pr": ""
        },
        {
            "id": "5",
            "stage": "validate_ticket",
            "data_type": "currency_formatting",
            "issue": "Currency notation discrepancies ('Rs. 3,75,000/-' vs 'INR 375000') flagged as amount mismatch",
            "reason": "Validator treats text formatting difference as conflicting budget amounts instead of recognizing monetary equivalence.",
            "test_case": "case_demo_currency_format",
            "expected_behavior": "Validator must parse numeric value and recognize that 'Rs. 3,75,000/-' equals 'INR 375000'.",
            "status": "open",
            "attempts": "0",
            "resolution_note": "",
            "linked_pr": ""
        },
        {
            "id": "6",
            "stage": "validate_ticket",
            "data_type": "same_day_dates",
            "issue": "Same-day administrative sanction flagged as suspicious chronological ordering",
            "reason": "Validator strictly expects sanction date strictly before application date without permitting same-day administrative sanction.",
            "test_case": "case_demo_date_order",
            "expected_behavior": "Validator must allow same-day sanction order and application submission.",
            "status": "open",
            "attempts": "0",
            "resolution_note": "",
            "linked_pr": ""
        }
    ]

    fieldnames = [
        "id", "stage", "data_type", "issue", "reason", "test_case",
        "expected_behavior", "status", "attempts", "resolution_note", "linked_pr"
    ]
    with open(issues_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    df = pd.DataFrame(rows)
    df.to_excel(issues_xlsx, index=False)
    print(f"[OK] Saved {len(rows)} issues to {issues_csv} and {issues_xlsx}")


if __name__ == "__main__":
    generate_currency_problem()
    generate_date_order_problem()
    update_excel_and_csv()
